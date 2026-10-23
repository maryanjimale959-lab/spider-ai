"""DOCX parsing with python-docx.

Word files carry no page coordinates in a portable way, so passages are anchored to
their rendered paragraph element (``coordinate_source = dom_element``). That is stated
in the implementation notes instead of pretending page numbers exist.
"""
from __future__ import annotations

from pathlib import Path

import docx

from app.core.ids import new_id
from app.parsing.base import ParsedDocument, ParsedPassage, ParsedSection, clean_text


def _kind_for(paragraph: docx.text.paragraph.Paragraph) -> tuple[str, int]:
    style = (paragraph.style.name or "").strip() if paragraph.style else ""
    low = style.lower()
    if low.startswith("title"):
        return "title", 0
    if low.startswith("subtitle"):
        return "heading", 1
    if low.startswith("heading"):
        try:
            level = int(style.split()[-1])
        except ValueError:
            level = 1
        return "heading", level
    if "list" in low:
        return "list_item", 9
    if low in ("caption", "quote", "intense quote"):
        return "caption", 9
    return "paragraph", 9


def _split_long(paragraph_text: str, limit: int = 1400) -> list[str]:
    text = clean_text(paragraph_text)
    if len(text) <= limit:
        return [text] if text else []
    sentences: list[str] = []
    buffer = ""
    for chunk in (text + " ").replace(". ", ".\x01").replace("? ", "?\x01").replace("! ", "!\x01").split("\x01"):
            buffer += chunk
            if len(buffer) >= limit:
                sentences.append(buffer.strip())
                buffer = ""
    if buffer.strip():
        sentences.append(buffer.strip())
    return sentences or [text[:limit]]


def parse_docx(path: Path) -> ParsedDocument:
    doc = ParsedDocument(format="docx", parser="python-docx", positioning_quality="dom_element")
    word = docx.Document(str(path))
    sort_index = 0
    section: ParsedSection | None = None

    for paragraph in word.paragraphs:
        texts = _split_long(paragraph.text)
        for text in texts:
            kind, level = _kind_for(paragraph)
            if kind in ("title", "heading"):
                section = ParsedSection(id=new_id(), index=len(doc.sections), level=level, title=text[:300], kind=kind)
                doc.sections.append(section)
                if kind == "title" and not doc.title:
                    doc.title = text[:200]
            doc.passages.append(
                ParsedPassage(id=new_id(), sort_index=sort_index, kind=kind, text=text,
                              section_id=section.id if section else None, coordinate_source="dom_element")
            )
            sort_index += 1

    for table in word.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            text = " | ".join(c for c in cells if c)
            if not text:
                continue
            doc.passages.append(
                ParsedPassage(id=new_id(), sort_index=sort_index, kind="table_row", text=text[:2000],
                              section_id=section.id if section else None, coordinate_source="dom_element",
                              meta={"cells": cells})
            )
            sort_index += 1

    if not doc.title:
        core_title = (word.core_properties.title or "").strip()
        doc.title = core_title or path.stem
    doc.notes.append(
        "DOCX has no stable page coordinates in the file; passages are anchored to their rendered paragraph, "
        "so the spider stands on the exact paragraph but page numbers are not reported."
    )
    return doc.finalize()
